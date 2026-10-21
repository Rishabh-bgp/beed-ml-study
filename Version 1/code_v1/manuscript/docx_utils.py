"""Helpers to build conference-compliant .docx files with python-docx."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Inches, RGBColor

FONT = "Times New Roman"


class Refs:
    """Numbered citations in order of first appearance."""
    def __init__(self, bib):
        self.bib, self.order = bib, []

    def cite(self, *keys):
        nums = []
        for k in keys:
            if k not in self.bib:
                raise KeyError(k)
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        nums = sorted(set(nums))
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            parts.append(str(nums[i]) if j == i else (f"{nums[i]},{nums[j]}" if j == i + 1 else f"{nums[i]}–{nums[j]}"))
            i = j + 1
        return "[" + ",".join(parts) + "]"


def _set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge, val in kw.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        for k, v in val.items():
            el.set(qn(f"w:{k}"), str(v))


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    run.font.name = FONT; run.font.size = Pt(10)


class Doc:
    def __init__(self, line_spacing=2.0, space_after=0):
        self.d = Document()
        self.ls = line_spacing
        self.fig_n = 0; self.tab_n = 0
        self.fig_prefix = ""; self.tab_prefix = ""
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, m, Cm(2.5))
        add_page_number_footer(sec)
        st = self.d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = line_spacing; pf.space_after = Pt(space_after); pf.space_before = Pt(0)
        for lvl, size in [(1, 14), (2, 12)]:
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(0)
            h.paragraph_format.line_spacing = line_spacing
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.keep_with_next = True
        t = self.d.styles["Title"]
        t.font.name = FONT; t.font.size = Pt(16); t.font.bold = True; t.font.color.rgb = RGBColor(0, 0, 0)
        trpr = t.element.get_or_add_rPr(); trf = trpr.find(qn("w:rFonts"))
        if trf is None:
            trf = OxmlElement("w:rFonts"); trpr.append(trf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            trf.set(qn(a), FONT)
        t.paragraph_format.line_spacing = line_spacing
        # remove the Title style's bottom border
        pPr = t.element.get_or_add_pPr()
        for b in pPr.findall(qn("w:pBdr")):
            pPr.remove(b)

    # ---- text -------------------------------------------------------------------------------
    def title(self, text):
        p = self.d.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(p, text)
        return p

    def h1(self, text):
        return self.d.add_heading(text, level=1)

    def h2(self, text):
        return self.d.add_heading(text, level=2)

    def para(self, text, align="justify", italic=False, size=None, bold_prefix=None, center=False):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (WD_ALIGN_PARAGRAPH.JUSTIFY if align == "justify" else WD_ALIGN_PARAGRAPH.LEFT)
        if bold_prefix:
            r = p.add_run(bold_prefix + " "); r.bold = True
            if size: r.font.size = Pt(size)
        self._runs(p, text, italic=italic, size=size)
        return p

    def runin(self, head, text):
        """Third-level heading: bold run-in, unnumbered, text continues on the same line."""
        return self.para(text, bold_prefix=head.rstrip(".") + ".")

    def bullets(self, items):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = self.ls
            self._runs(p, it)

    def _runs(self, p, text, italic=False, size=None):
        """Minimal inline markup: **bold**, *italic*, ^{sup}, _{sub}."""
        tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\})", text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**"):
                r = p.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith("*"):
                r = p.add_run(tok[1:-1]); r.italic = True
            elif tok.startswith("^{"):
                r = p.add_run(tok[2:-1]); r.font.superscript = True
            elif tok.startswith("_{"):
                r = p.add_run(tok[2:-1]); r.font.subscript = True
            else:
                r = p.add_run(tok)
            if italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)

    # ---- floats -----------------------------------------------------------------------------
    def figure(self, path, title, legend="", width_in=6.2):
        self.fig_n += 1
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width_in))
        cap = self.d.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = cap.add_run(f"Figure {self.fig_prefix}{self.fig_n}. {title}" + ("" if title.endswith(".") else ".")); r.bold = True
        if legend:
            cap.add_run(" ")
            self._runs(cap, legend)
        assert len(title.split()) <= 15, title
        assert len(legend.split()) <= 300
        return self.fig_n

    def table(self, title, header, rows, note="", col_widths=None, font_size=9):
        self.tab_n += 1
        assert len(title.split()) <= 15, title
        cap = self.d.add_paragraph(); cap.paragraph_format.keep_with_next = True
        r = cap.add_run(f"Table {self.tab_prefix}{self.tab_n}. {title}"); r.bold = True
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            self._cell(t.cell(0, j), h, font_size, bold=True)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                self._cell(t.cell(i + 1, j), v, font_size)
        line = {"val": "single", "sz": 6, "color": "000000"}
        none = {"val": "nil"}
        for i, row in enumerate(t.rows):
            for c in row.cells:
                _set_cell_border(c, top=line if i in (0, 1) else none, bottom=line if i in (0, len(t.rows) - 1) else none,
                                 left=none, right=none)
        if col_widths:
            t.autofit = False
            for j, w in enumerate(col_widths):
                t.columns[j].width = Cm(w)
                for row in t.rows:
                    row.cells[j].width = Cm(w)
        if note:
            n = self.d.add_paragraph(); n.paragraph_format.line_spacing = 1.5
            self._runs(n, note, size=10)
            assert len(note.split()) <= 300
        else:
            self.d.add_paragraph().paragraph_format.line_spacing = 1.0
        return self.tab_n

    def _cell(self, cell, text, size, bold=False):
        cell.text = ""
        p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._runs(p, str(text), size=size)
        if bold:
            for r in p.runs:
                r.bold = True

    def page_break(self):
        self.d.add_page_break()

    def save(self, path, title=""):
        import datetime as _dt
        cp = self.d.core_properties
        cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.title = title
        cp.created = cp.modified = _dt.datetime(2026, 10, 8)
        self.d.save(str(path))
